from io import BytesIO
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from uuid import UUID
from app.api.v1.dependencies import CurrentUser, require_auth
from app.core.database import get_db
from app.schemas import ResponderCreate, ResponderDetailsResponse, ResponderListItem, ResponderAdminUpdate
from app.schemas import ResponderDeploymentCreate, ResponderDeploymentItem, ResponderDeploymentsResponse
from app.schemas import ResponderActiveSiteResponse, ResponderAdminActiveSiteUpdate
from app.services import responder_service, responder_deployment_service, responder_app_service

router = APIRouter(prefix="/responders", tags=["responders"])


@router.get("/all", dependencies=[Depends(require_auth)], response_model=list[ResponderListItem])
async def get_all_responders(db:AsyncSession = Depends(get_db)) -> list[ResponderListItem]:
    return await responder_service.get_all_responders(db=db)


@router.get("/additional-details/{responder_id}", dependencies=[Depends(require_auth)], response_model=ResponderDetailsResponse)
async def get_responder_details(responder_id: str, db:AsyncSession = Depends(get_db)) -> ResponderDetailsResponse:
    return await responder_service.get_responder_details(responder_id=responder_id, db=db)


@router.get("/template", dependencies=[Depends(require_auth)])
async def download_responder_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Responders"

    headers = ["First Name", "Last Name", "Phone Number", "Location ID"]
    widths = [20, 20, 25, 12]

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 25
    ws.sheet_view.pane = None
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=responder_template.xlsx"},
    )


@router.post("/bulk", response_model=list[ResponderListItem])
async def bulk_create_responders(
    responders: list[ResponderCreate],
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(require_auth)) -> list[ResponderListItem]:

    return await responder_service.bulk_create_responders(responders=responders, db=db, user_id=user.id)


# Phone changes do not invalidate the responder JWT or push subscriptions; both key off responder_id.
@router.patch("/{responder_id}", dependencies=[Depends(require_auth)], response_model=ResponderListItem)
async def update_responder(
    responder_id: UUID,
    payload: ResponderAdminUpdate,
    db: AsyncSession = Depends(get_db),
) -> ResponderListItem:
    return await responder_service.update_responder_details(responder_id=responder_id, payload=payload, db=db)


# --- Deployments (admin-managed site membership) ---

@router.get("/{responder_id}/deployments", dependencies=[Depends(require_auth)], response_model=ResponderDeploymentsResponse)
async def list_responder_deployments(
    responder_id: UUID,
    db: AsyncSession = Depends(get_db),
) -> ResponderDeploymentsResponse:
    return await responder_deployment_service.list_deployments(db=db, responder_id=responder_id)


@router.post("/{responder_id}/deployments", response_model=ResponderDeploymentItem, status_code=201)
async def open_responder_deployment(
    responder_id: UUID,
    payload: ResponderDeploymentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: CurrentUser = Depends(require_auth),
) -> ResponderDeploymentItem:
    return await responder_deployment_service.open_deployment(
        db=db, responder_id=responder_id, payload=payload, current_user=current_user
    )


@router.post("/{responder_id}/deployments/{deployment_id}/close", response_model=ResponderDeploymentItem)
async def close_responder_deployment(
    responder_id: UUID,
    deployment_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: CurrentUser = Depends(require_auth),
) -> ResponderDeploymentItem:
    return await responder_deployment_service.close_deployment(
        db=db, responder_id=responder_id, deployment_id=deployment_id, current_user=current_user
    )


# Admin presence override: set a responder's active site (must still be a
# site they're eligible for) or clear it (null → falls back to home site).
@router.put("/{responder_id}/active-site", response_model=ResponderActiveSiteResponse)
async def admin_set_active_site(
    responder_id: UUID,
    payload: ResponderAdminActiveSiteUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: CurrentUser = Depends(require_auth),
) -> ResponderActiveSiteResponse:
    return await responder_app_service.admin_set_active_site(
        db=db, responder_id=responder_id, location_id=payload.location_id,
        current_user=current_user,
    )
